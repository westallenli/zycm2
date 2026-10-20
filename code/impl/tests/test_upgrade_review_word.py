"""Anonymous P13 Word preview/recovery: no overwrite, conflict or parser bypass."""
import base64
from copy import deepcopy
from dataclasses import asdict, replace
from http.client import HTTPConnection
import io
import json
import threading
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from panda_brain import review_exports as exports, review_word as word, review_reports
from panda_brain import review_meeting
from panda_brain.panel import make_server
from panda_brain.persistence import load_workspace, save_workspace
from panda_brain.review_domain import ReviewLedger
from tests import test_upgrade_review_reports as fixture


def edited(raw, text='匿名线下补充：仍待核实，不自动定责。'):
    document = Document(io.BytesIO(raw))
    document.add_paragraph(text)
    stream = io.BytesIO(); document.save(stream)
    return stream.getvalue()


def package_change(raw, name, content):
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(result, 'w', zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item.filename, content if item.filename == name else source.read(item))
    return result.getvalue()


class WordTests(unittest.TestCase):
    def setUp(self):
        fixture.ReportSnapshotTests.setUp(self)
        self.rid = self.state.handle_report_generate()
        self.state.handle_report_edit(self.rid, self.state.review_ledger.reports[self.rid].content_md)
        self.state.handle_report_status(self.rid, 'pending_review')
        self.state.handle_report_status(self.rid, 'published')
        self.original = self.state.review_ledger.reports[self.rid]
        self.raw = edited(exports.export_version(self.original, 'docx'))

    def preview(self, **changes):
        values = dict(token=self.tokens[1], report_id=self.rid,
                      expected_version=self.state.workspace.version, raw=self.raw)
        values.update(changes)
        return word.preview(self.state, **values)

    def test_preview_is_readonly_and_confirm_creates_only_new_pending_report(self):
        before = (deepcopy(self.state.review_ledger.issues), deepcopy(self.state.review_ledger.rects),
                  deepcopy(self.state.review_ledger.hypotheses), deepcopy(self.state.review_ledger.experiences))
        version = self.state.workspace.version
        files = (self.state.db_path.read_bytes(), self.state.reviews_path.read_bytes())
        result = self.preview()
        self.assertIn('匿名线下补充', result['content_md'])
        self.assertEqual(version, self.state.workspace.version)
        self.assertEqual(files, (self.state.db_path.read_bytes(), self.state.reviews_path.read_bytes()))
        rid = word.confirm(self.state, self.tokens[1], result['preview_id'], version)
        recovered = self.state.review_ledger.reports[rid]
        self.assertEqual('pending_review', recovered.status)
        self.assertEqual(2, recovered.version)
        self.assertEqual(self.original, self.state.review_ledger.reports[self.rid])
        self.assertEqual(before, (self.state.review_ledger.issues, self.state.review_ledger.rects,
                                 self.state.review_ledger.hypotheses, self.state.review_ledger.experiences))
        self.assertEqual('', review_reports.freshness_error(self.state, recovered))
        self.assertEqual(self.rid, recovered.imported_from_report_id)
        self.assertEqual(recovered, ReviewLedger.load(self.state.reviews_path).reports[rid])
        self.assertFalse(load_workspace(self.state.db_path).verify_audit_chain())
        self.assertEqual(rid, word.confirm(self.state, self.tokens[1], result['preview_id'], version))
        self.assertEqual(2, len(self.state.review_ledger.reports))
        self.state.handle_report_status(rid, 'published')
        self.assertEqual('published', self.state.review_ledger.reports[rid].status)

    def test_new_report_or_business_record_after_preview_blocks_import(self):
        for change in ('report', 'meeting'):
            with self.subTest(change=change):
                result = self.preview(); version = self.state.workspace.version
                if change == 'report':
                    self.state.handle_report_generate()
                else:
                    fixture.ReportSnapshotTests.meeting(self)
                before = self.state.reviews_path.read_bytes()
                with self.assertRaises(ValueError):
                    word.confirm(self.state, self.tokens[1], result['preview_id'], version)
                self.assertEqual(before, self.state.reviews_path.read_bytes())
                if change == 'report':
                    # Next preview uses the latest report so the second branch starts fresh.
                    self.rid = max(self.state.review_ledger.reports.values(), key=lambda row:row.version).report_id
                    self.raw = edited(exports.export_version(self.state.review_ledger.reports[self.rid], 'docx'))

    def test_owner_session_event_expiry_and_original_identity_are_enforced(self):
        with self.assertRaises(PermissionError):
            self.preview(token=self.tokens[0])
        result = self.preview(); version = self.state.workspace.version
        new_token, _ = self.state.member_login('13912340002', '123456')
        with self.assertRaises(ValueError):
            word.confirm(self.state, new_token, result['preview_id'], version)
        with patch('panda_brain.review_word.time.monotonic', return_value=10**12):
            with self.assertRaises(ValueError):
                word.confirm(self.state, self.tokens[1], result['preview_id'], version)
        self.state.review_ledger.reports[self.rid] = replace(self.original, content_md='后来改过的正文')
        with self.assertRaisesRegex(ValueError, '版本不匹配'):
            self.preview()
        self.state.active_event_id = 'fire-other'
        with self.assertRaises((PermissionError, ValueError)):
            word.confirm(self.state, self.tokens[1], result['preview_id'], version)

    def test_chinese_simple_tables_and_body_only_are_preserved(self):
        report = replace(self.original, content_md='## 匿名正文\n文字 <script>literal</script>\n'
                         '| 项目 | 记录 |\n| --- | --- |\n| 空值 |  |\n| 依据 | 甲\\|乙 |')
        raw = exports.export_version(report, 'docx')
        document = Document(io.BytesIO(raw)); document.paragraphs[0].text = '线下标题被改'
        # WPS removes dc:identifier on save; the body version marker must still match.
        document.core_properties.identifier = ''
        document.add_paragraph('正文补充')
        stream = io.BytesIO(); document.save(stream)
        text = word.read_body(stream.getvalue(), report)
        self.assertIn('## 匿名正文', text)
        self.assertIn('甲\\|乙', text)
        self.assertIn('| 空值 |  |', text)
        self.assertIn('<script>literal</script>', text)
        self.assertIn('正文补充', text)
        self.assertNotIn('线下标题被改', text)

    def test_missing_marker_identifier_revision_and_external_link_are_rejected(self):
        for change in ('marker', 'identity', 'version_marker', 'revision', 'image', 'link'):
            with self.subTest(change=change):
                document = Document(io.BytesIO(self.raw))
                if change == 'identity':document.core_properties.identifier = 'other'
                elif change == 'version_marker':
                    document.core_properties.identifier = ''
                    for row in document.paragraphs:
                        if row.text == word.marker_text(self.original): row.text += ' changed'
                elif change == 'marker':
                    for row in document.element.body.iter(qn('w:bookmarkStart')):
                        if row.get(qn('w:name')) == word.BODY_BOOKMARK: row.set(qn('w:name'), 'removed')
                elif change in {'revision', 'image'}:
                    document.paragraphs[-1]._p.append(OxmlElement('w:ins' if change == 'revision' else 'w:drawing'))
                elif change == 'link':
                    document.part.relate_to('https://example.invalid/private',
                        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
                stream=io.BytesIO();document.save(stream)
                with self.assertRaises(ValueError):self.preview(raw=stream.getvalue())

    def test_bad_zip_entities_paths_and_size_are_rejected(self):
        for raw in (b'not-docx', b'x'*(word.MAX_FILE_BYTES+1),
                    package_change(self.raw, 'word/document.xml', b'<!DOCTYPE x [<!ENTITY y "x">]><x>&y;</x>'),
                    package_change(self.raw, 'word/document.xml', '<!DOCTYPE x><x/>'.encode('utf-16'))):
            with self.subTest(size=len(raw)):
                with self.assertRaises(ValueError):self.preview(raw=raw)
        package = io.BytesIO()
        with zipfile.ZipFile(package, 'w') as archive:archive.writestr('../outside.xml', '<x/>')
        with self.assertRaises(ValueError):self.preview(raw=package.getvalue())

    def test_failed_ledger_save_restores_audit_and_allows_retry(self):
        snapshot = self.state.db_path.parent/'older-snapshot.db'
        older = deepcopy(self.state.workspace); older.scenario_id = 'anonymous-older-snapshot'
        save_workspace(older, snapshot)
        self.state.registry.events['fire-a'] = replace(
            self.state.registry.events['fire-a'], snapshot_path=str(snapshot))
        result=self.preview();version=self.state.workspace.version
        original_events=tuple(self.state.workspace.events); before=self.state.reviews_path.read_bytes()
        with patch.object(ReviewLedger, 'save', side_effect=OSError('anonymous save failure')):
            with self.assertRaises(OSError):word.confirm(self.state,self.tokens[1],result['preview_id'],version)
        self.assertEqual(before,self.state.reviews_path.read_bytes())
        self.assertEqual(original_events,tuple(self.state.workspace.events))
        loaded = load_workspace(self.state.db_path)
        normalize = lambda rows: json.loads(json.dumps([asdict(row) for row in rows]))
        self.assertEqual(normalize(original_events), normalize(loaded.events))
        self.assertFalse(loaded.verify_audit_chain())
        self.assertEqual('anonymous-older-snapshot', load_workspace(snapshot).scenario_id)
        self.assertEqual(1,len(self.state.review_ledger.reports))
        self.assertTrue(word.confirm(self.state,self.tokens[1],result['preview_id'],version))


class WordHTTPTests(unittest.TestCase):
    def setUp(self):
        fixture.ReportSnapshotTests.setUp(self)
        self.rid=self.state.handle_report_generate()
        self.raw=edited(exports.export_version(self.state.review_ledger.reports[self.rid],'docx'))
        self.server=make_server(self.state,port=0)
        threading.Thread(target=self.server.serve_forever,daemon=True).start()
        self.addCleanup(self.server.server_close);self.addCleanup(self.server.shutdown)

    def fields(self):
        return dict(event_id='fire-a',csrf=review_meeting.csrf(self.state,self.tokens[1]),
                    expected_version=str(self.state.workspace.version))

    def request(self, path, fields, cookie=None, length=None):
        conn=HTTPConnection('127.0.0.1',self.server.server_port,timeout=10)
        raw=fields if isinstance(fields,str) else json.dumps(fields)
        headers={'Cookie':cookie or 'member_session='+self.tokens[1],'Content-Type':'application/json'}
        if length is not None:headers['Content-Length']=str(length);raw=''
        try:
            conn.request('POST',path,raw,headers);response=conn.getresponse()
            return response.status,json.loads(response.read())
        finally:conn.close()

    def test_preview_confirm_retry_auth_csrf_event_and_duplicate_fields(self):
        fields={**self.fields(),'report_id':self.rid,'file_b64':base64.b64encode(self.raw).decode()}
        for cookie in ('none=none','member_session='+self.tokens[0],'panel_session='+self.state.login('anonymous-report-secret')):
            self.assertEqual(403,self.request('/review/word-preview',fields,cookie)[0])
        self.assertEqual(403,self.request('/review/word-preview',{**fields,'csrf':'wrong'})[0])
        self.assertEqual(403,self.request('/review/word-preview',{**fields,'event_id':'fire-other'})[0])
        duplicate=json.dumps(fields)[:-1]+',"report_id":"other"}'
        self.assertEqual(409,self.request('/review/word-preview',duplicate)[0])
        self.assertEqual(413,self.request('/review/word-preview',{},length=word.MAX_REQUEST_BYTES+1)[0])
        self.assertEqual(413,self.request('/review/word-preview',{},length=-1)[0])
        status,result=self.request('/review/word-preview',fields)
        self.assertEqual(200,status)
        confirm={**self.fields(),'preview_id':result['preview_id'],'confirm':'1'}
        status,done=self.request('/review/word-import',confirm);self.assertEqual(200,status)
        self.assertEqual(done,self.request('/review/word-import',confirm)[1])
        self.assertEqual(2,len(self.state.review_ledger.reports))
        self.assertEqual('pending_review',self.state.review_ledger.reports[done['report_id']].status)


if __name__ == '__main__':unittest.main()
