"""P13-03: real versioned documents, immutable provenance and HTTP scope."""
import io
import json
import threading
import unittest
from dataclasses import replace
from http.client import HTTPConnection
from unittest.mock import patch
from urllib.parse import urlencode

from panda_brain import review_exports, review_reports
from panda_brain.panel import make_server
from panda_brain.review_domain import ReviewReport, RectTask
from panda_brain.reviewview import render_review_page
from tests import test_upgrade_review_reports as fixture


class DocumentDeliveryTests(unittest.TestCase):
    def test_editable_word_and_pdf_preserve_chinese_metadata_tables_and_empty_values(self):
        from docx import Document
        from pypdf import PdfReader
        report = ReviewReport('report-anon', 'fire-anon', 3, status='published',
                              created_at='2026-10-07T00:00:00+00:00', reviewed_by='匿名负责人',
                              reviewed_at='2026-10-07T01:00:00+00:00',
                              content_md='## 中文正文\n原始事实待核；不得自动定责。\n\n'
                              '| 项目 | 记录 |\n| --- | --- |\n| 空值 |  |\n'
                              '| 依据 | 中文甲乙丙 <script>literal</script> |')
        docx = Document(io.BytesIO(review_exports.export_version(report, 'docx')))
        text = '\n'.join(p.text for p in docx.paragraphs)
        for value in ('报告版本：v3', '已发布', '匿名负责人', report.created_at, '原始事实待核'):
            self.assertIn(value, text)
        self.assertEqual('', docx.tables[0].cell(1, 1).text)
        self.assertIn('<script>literal</script>', docx.tables[0].cell(2, 1).text)
        docx.add_paragraph('可编辑新增')
        stream = io.BytesIO()
        docx.save(stream)
        self.assertEqual('可编辑新增', Document(io.BytesIO(stream.getvalue())).paragraphs[-1].text)
        pdf = PdfReader(io.BytesIO(review_exports.export_version(report, 'pdf')))
        extracted = ''.join(page.extract_text() for page in pdf.pages)
        for value in ('报告版本：v3', '匿名负责人', '中文甲乙丙', '<script>literal</script>', '空值', '第 1 页'):
            self.assertIn(value, extracted)

    def test_long_chinese_table_pages_keep_every_row_and_blank_body_is_honest(self):
        from pypdf import PdfReader
        rows = '\n'.join(f'| 行{i:03} | ' + '匿名长记录待核。'*25 + ' |' for i in range(36))
        report = ReviewReport('long', 'fire-anon', 1, content_md='## 长表格\n| 序号 | 内容 |\n| --- | --- |\n'+rows)
        pdf = PdfReader(io.BytesIO(review_exports.export_version(report, 'pdf')))
        text = ''.join(page.extract_text() for page in pdf.pages)
        self.assertGreater(len(pdf.pages), 2)
        for page in pdf.pages:
            self.assertLessEqual(page.extract_text().count('序号'), 1)
        for i in range(36):
            self.assertIn(f'行{i:03}', text)
        oversized = replace(report, content_md='| 项目 | 内容 |\n| --- | --- |\n| 超长 | '+ '匿名待核。'*3000+' |')
        with self.assertRaisesRegex(ValueError, '拆分长表格单元格'):
            review_exports.export_version(oversized, 'pdf')
        self.assertIn('历史版本未保存来源快照', review_exports.report_markdown(replace(report, content_md='')))
        self.assertTrue(review_exports.export_version(replace(report, content_md=''), 'docx').startswith(b'PK'))
        with patch.dict('os.environ', {'PANDA_REPORT_FONT': 'missing-anonymous-font.ttf'}):
            with self.assertRaisesRegex(ValueError, '字体未配置'):
                review_exports.export_version(report, 'pdf')


class VersionHTTPTests(unittest.TestCase):
    def setUp(self):
        fixture.ReportSnapshotTests.setUp(self)
        self.state.review_ledger.add_rect(RectTask('rect-a', ('issue-a',), '生成时措施', '模拟', '匿名0',
                                                  event_id='fire-a', status='in_progress'))
        self.rid = self.state.handle_report_generate()
        self.original = self.state.review_ledger.reports[self.rid]
        self.server = make_server(self.state, port=0)
        threading.Thread(target=self.server.serve_forever, daemon=True).start()
        self.addCleanup(self.server.server_close)
        self.addCleanup(self.server.shutdown)

    def get(self, path, cookie=None):
        conn = HTTPConnection('127.0.0.1', self.server.server_port, timeout=15)
        try:
            conn.request('GET', path, headers={'Cookie': cookie or 'member_session='+self.tokens[1]})
            response = conn.getresponse()
            return response.status, dict(response.getheaders()), response.read()
        finally:
            conn.close()

    def url(self, route='/export/review', **changes):
        return route+'?'+urlencode(dict(event_id='fire-a', report_id=self.rid, **changes))

    def test_real_downloads_are_selected_version_and_read_only_after_new_business_records(self):
        self.state.review_ledger.rects['rect-a'] = replace(self.state.review_ledger.rects['rect-a'], measures='后续措施')
        fixture.ReportSnapshotTests.meeting(self)
        self.state.handle_report_generate()
        before = (self.state.workspace.version, tuple(self.state.workspace.events), self.state.reviews_path.read_bytes())
        for format, magic in (('md', b'#'), ('docx', b'PK'), ('pdf', b'%PDF-')):
            status, headers, body = self.get(self.url(format=format))
            self.assertEqual(200, status)
            self.assertTrue(body.startswith(magic))
            self.assertEqual(len(body), int(headers['Content-Length']))
            self.assertIn(f'-v1.{format}', headers['Content-Disposition'])
            if format == 'md':
                self.assertIn('生成时措施', body.decode())
                self.assertNotIn('后续措施', body.decode())
        self.assertEqual(before, (self.state.workspace.version, tuple(self.state.workspace.events), self.state.reviews_path.read_bytes()))
        self.assertEqual(self.original, self.state.review_ledger.reports[self.rid])

    def test_source_links_use_snapshot_and_current_rects_are_explicitly_separate(self):
        self.state.review_ledger.rects['rect-a'] = replace(self.state.review_ledger.rects['rect-a'], measures='后续措施')
        status, _, body = self.get(self.url('/review/report-source', ref='rect-a'))
        self.assertEqual(200, status)
        self.assertIn('生成时措施', body.decode())
        self.assertNotIn('后续措施', body.decode())
        self.assertEqual(409, self.get(self.url('/review/report-source', ref='missing-new-record'))[0])
        page = render_review_page(self.state, role='owner', tab='report', selected_report=self.rid)
        for value in ('当前整改状态（后续更新，不改写报告正文）', '后续措施', '生成时原始资料', '下载 Word', '下载 PDF'):
            self.assertIn(value, page)
        self.assertEqual(self.original, self.state.review_ledger.reports[self.rid])

    def test_unauthorized_cross_event_duplicate_and_invalid_format_are_rejected(self):
        commander = self.state.login('anonymous-report-secret')
        self.assertEqual(200, self.get(self.url(format='md'), 'panel_session='+commander)[0])
        for route in ('/export/review', '/review/report-source'):
            for cookie in ('invalid_session=none', 'member_session='+self.tokens[0]):
                self.assertEqual(404, self.get(self.url(route), cookie)[0])
            self.assertEqual(409, self.get(self.url(route)+'&event_id=fire-b')[0])
            self.assertEqual(409, self.get(route+'?event_id=fire-b&report_id='+self.rid)[0])
        self.state.review_ledger.reports['foreign'] = ReviewReport('foreign', 'fire-b', 1, content_md='他事件私密正文')
        self.assertEqual(409, self.get('/export/review?event_id=fire-a&report_id=foreign&format=md')[0])
        self.assertEqual(409, self.get(self.url(format='html'))[0])

    def test_corrupt_and_missing_history_never_fabricate_sources_dependencies_fail_clearly(self):
        self.state.review_ledger.reports[self.rid] = replace(self.original, source_snapshot_json='{}')
        self.assertEqual(409, self.get(self.url(format='pdf'))[0])
        self.assertEqual(409, self.get(self.url('/review/report-source'))[0])
        page = render_review_page(self.state, role='owner', tab='report', selected_report=self.rid)
        self.assertIn('快照校验失败', page)
        self.state.review_ledger.reports[self.rid] = replace(self.original, source_snapshot_json='', source_digest='')
        self.assertEqual(200, self.get(self.url(format='md'))[0])
        self.assertEqual(409, self.get(self.url('/review/report-source'))[0])
        self.state.review_ledger.reports[self.rid] = self.original
        with patch('panda_brain.review_exports._docx', side_effect=ImportError('anonymous missing dependency')):
            self.assertEqual(503, self.get(self.url(format='docx'))[0])
        with patch.dict('os.environ', {'PANDA_REPORT_FONT': 'missing-anonymous-font.ttf'}):
            status, _, body = self.get(self.url(format='pdf'))
            self.assertEqual(409, status)
            self.assertIn('PANDA_REPORT_FONT', body.decode())


if __name__ == '__main__':
    unittest.main()
