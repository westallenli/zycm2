"""Render one immutable report version; no business mutations or network IO."""
from __future__ import annotations

import io
import os
import re
import threading
from datetime import datetime, timezone
from pathlib import Path
from xml.sax.saxutils import escape

_PDF_LOCK = threading.RLock()
CONTENT_TYPES = {
    'md': 'text/markdown; charset=utf-8',
    'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    'pdf': 'application/pdf',
}


def report_markdown(report):
    """Metadata is authoritative even when the owner edits the body headings."""
    return ('# 破局智援参谋 火情复盘报告\n\n'
            f'- 报告编号：{report.report_id}\n- 事件编号：{report.event_id}\n'
            f'- 报告版本：v{report.version}\n- 生成时间：{report.created_at or "未记录"}\n'
            f'- 审核状态：{report.status_word()}\n'
            f'- 报告范围：{"阶段复盘（非最终结案）" if report.stage_report else "事件复盘"}\n'
            f'- 审核发布人：{report.reviewed_by or "尚未发布"}\n'
            f'- 审核时间：{report.reviewed_at or "未记录"}\n'
            + (f'- 来源记录截至：第 {report.source_audit_sequence} 条（生成时资料另存核对）\n\n' if report.source_snapshot_json else '- 历史版本未保存来源快照\n\n')
            + '以下为所选版本正文，生成时状态和来源时点按原文保留。\n\n'
            + report.content_md)


def _cells(line):
    return [cell.strip().replace('\\|', '|') for cell in re.split(r'(?<!\\)\|', line.strip().strip('|'))]


def blocks(text):
    """Deterministic Markdown subset; unsupported syntax remains readable text."""
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if (index + 1 < len(lines) and '|' in line
                and all(re.fullmatch(r':?-{3,}:?', cell) for cell in _cells(lines[index + 1]))):
            rows = [_cells(line)]
            width = len(rows[0])
            index += 2
            while index < len(lines) and '|' in lines[index] and len(_cells(lines[index])) == width:
                rows.append(_cells(lines[index]))
                index += 1
            yield 'table', rows
            continue
        heading = re.match(r'^(#{1,6})\s+(.+)$', line)
        if heading:
            yield 'heading', (len(heading[1]), heading[2])
        elif line.strip():
            yield 'text', line
        else:
            yield 'space', ''
        index += 1


def _plain(text):
    return text.replace('**', '').replace('`', '')


def _docx(text, report=None):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.left_margin = section.right_margin = Inches(.8)
    section.top_margin = section.bottom_margin = Inches(.7)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        style = document.styles[name]
        style.font.name = 'SimSun'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'SimSun')
    document.styles['Normal'].font.size = Pt(10.5)
    document.styles['Normal'].paragraph_format.space_after = Pt(5)
    document.core_properties.author = '破局智援参谋'
    document.core_properties.title = '破局智援参谋 火情复盘报告'
    document.core_properties.created = document.core_properties.modified = datetime.now(timezone.utc)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('第 ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    footer.add_run(' 页')
    first_title = True
    for kind, value in blocks(text):
        if kind == 'heading':
            level, title = value
            if first_title:
                document.add_paragraph(_plain(title), 'Title')
                first_title = False
            else:
                document.add_heading(_plain(title), min(level, 3))
        elif kind == 'table':
            table = document.add_table(rows=0, cols=len(value[0]))
            table.style = 'Table Grid'
            for number, row in enumerate(value):
                cells = table.add_row().cells
                for cell, content in zip(cells, row):
                    cell.text = _plain(content)
                if number == 0:
                    repeat = OxmlElement('w:tblHeader')
                    table.rows[0]._tr.get_or_add_trPr().append(repeat)
                    for cell in cells:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
        elif kind == 'text':
            document.add_paragraph(_plain(value))
    if report is not None:
        from .review_word import identity, BODY_BOOKMARK, marker_text
        document.core_properties.identifier = identity(report)
        prefix = text[:len(text)-len(report.content_md)] if report.content_md else text
        count = sum(kind != 'space' for kind, _ in blocks(prefix))
        target = list(document.element.body)[count]
        marker = document.add_paragraph()
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), '100')
        start.set(qn('w:name'), BODY_BOOKMARK)
        marker._p.append(start)
        marker.add_run(marker_text(report))
        end = OxmlElement('w:bookmarkEnd')
        end.set(qn('w:id'), '100')
        marker._p.append(end)
        target.addprevious(marker._p)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _pdf(text):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont, TTFError
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.platypus.doctemplate import LayoutError
    font_path = os.environ.get('PANDA_REPORT_FONT', '')
    if not font_path:
        font_path = str(Path(os.environ.get('WINDIR', 'C:/Windows'))/'Fonts'/'simsun.ttc')
    if not Path(font_path).is_file():
        raise ValueError('PDF 中文字体未配置，请设置 PANDA_REPORT_FONT 为允许嵌入的中文 TrueType 字体。')
    with _PDF_LOCK:
        try:
            font = TTFont('PandaReportChinese', font_path, subfontIndex=0)
        except (TTFError, OSError) as exc:
            raise ValueError('PDF 中文字体无法读取，请核对 PANDA_REPORT_FONT。') from exc
        missing = {char for char in text if not char.isspace() and ord(char) not in font.face.charToGlyph}
        if missing:
            raise ValueError('PDF 字体缺少报告所需字符，请配置覆盖这些字符的中文字体。')
        pdfmetrics.registerFont(font)
        body = ParagraphStyle('ReportBody', fontName=font.fontName, fontSize=10.5, leading=16,
                              wordWrap='CJK', splitLongWords=True, spaceAfter=5, alignment=TA_LEFT)
        heading = ParagraphStyle('ReportHeading', parent=body, fontSize=13, leading=19,
                                 spaceBefore=10, keepWithNext=True)
        title = ParagraphStyle('ReportTitle', parent=heading, fontSize=18, leading=26)
        def paragraph(value, style=body):
            return Paragraph(escape(_plain(value)).replace('\t', '    '), style)
        story = []
        first_title = True
        width = A4[0] - 100
        for kind, value in blocks(text):
            if kind == 'heading':
                story.append(paragraph(value[1], title if first_title else heading))
                first_title = False
            elif kind == 'table':
                table = Table([[paragraph(cell) for cell in row] for row in value],
                              colWidths=[width/len(value[0])]*len(value[0]),
                              repeatRows=1, splitByRow=1, splitInRow=0)
                table.setStyle(TableStyle([
                    ('GRID', (0, 0), (-1, -1), .4, colors.grey),
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#eeeeee')),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('LEFTPADDING', (0, 0), (-1, -1), 5),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 5),
                ]))
                story.append(table)
            elif kind == 'text':
                story.append(paragraph(value))
            else:
                story.append(Spacer(1, 3))
        stream = io.BytesIO()
        def footer(canvas, doc):
            canvas.setFont(font.fontName, 9)
            canvas.drawRightString(A4[0]-50, 27, f'第 {doc.page} 页')
        try:
            SimpleDocTemplate(stream, pagesize=A4, rightMargin=50, leftMargin=50,
                              topMargin=40, bottomMargin=45, title='破局智援参谋 火情复盘报告',
                              author='破局智援参谋').build(story, onFirstPage=footer, onLaterPages=footer)
        except LayoutError as exc:
            raise ValueError('报告表格或段落超出版面，请拆分长表格单元格后重新审核。') from exc
        return stream.getvalue()


def export_version(report, format):
    if format not in CONTENT_TYPES:
        raise ValueError('不支持的报告格式。')
    if len(report.content_md) > 400000:
        raise ValueError('报告正文过长，请拆分后审核。')
    text = report_markdown(report)
    return text.encode('utf-8') if format == 'md' else _docx(text, report) if format == 'docx' else _pdf(text)
