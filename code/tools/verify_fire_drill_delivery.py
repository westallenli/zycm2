"""Read delivery state, export its existing report, and compare preserved roots."""
from pathlib import Path
import hashlib,json,sys
W=Path(__file__).resolve().parents[1];N=W.parent
E=N/'docs/开发/灭火场景恢复与性能证据_20261017'
R=Path((E/'delivery-root.txt').read_text(encoding='utf-8'))
sys.dont_write_bytecode=True;sys.path.insert(0,str(R/'source'))
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
before=json.loads((E/'before.json').read_text(encoding='utf-8'))
preserved={}
for name,entry in before.items():
    root=Path(entry['root'])
    assert all(sha(root/'data'/p)==digest for p,digest in entry['data'].items()),name
    assert sha(root/'source-manifest.json')==entry['manifest_sha256'],name
    preserved[name]=dict(root=str(root),data_files=len(entry['data']),unchanged=True)
manifest=json.loads((R/'source-manifest.json').read_text(encoding='utf-8'))
assert all(sha(R/'source'/p)==digest for p,digest in manifest.items())
assert all(p=='config/settings.yaml' or sha(W/'impl'/p)==digest for p,digest in manifest.items())
package=json.loads((E/'package.json').read_text(encoding='utf-8'))
assert all(sha(R/'data'/p)==digest for p,digest in package['data_hashes'].items())
from panda_brain.resource_demo import open_state
from panda_brain import review_exports,review_reports,nine_flow
from panda_brain.read_scope import read_scope
state=open_state(R/'data')
with state._event_authority_lock,read_scope():
    assert not state.workspace.verify_audit_chain()
    report=next(iter(state.review_ledger.reports.values()))
    assert not review_reports.freshness_error(state,report)
    steps=nine_flow.summary(state)
    assert all(not step['stale'] for step in steps)
out=E/'artifacts';out.mkdir(exist_ok=True)
for fmt in ('md','docx','pdf'):
    (out/('北山林火阶段复盘.'+fmt)).write_bytes(review_exports.export_version(report,fmt))
from docx import Document
import fitz
document=Document(out/'北山林火阶段复盘.docx')
doc_text='\n'.join(p.text for p in document.paragraphs)+'\n'+'\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
pdf=fitz.open(out/'北山林火阶段复盘.pdf');pdf_text='\n'.join(p.get_text() for p in pdf)
for needle in ('青岚县北山','通道','第2版','照明','九步','阶段复盘'):
    assert needle in doc_text,('docx',needle)
    assert needle in pdf_text,('pdf',needle)
assert 'revision-' not in report.content_md
pdf[0].get_pixmap(matrix=fitz.Matrix(1.1,1.1)).save(out/'report-first-page.png')
result=dict(root=str(R),runtime=json.loads((R/'runtime.json').read_text(encoding='utf-8')),
    audit_valid=True,audit_count=state.workspace.version,source_files=len(manifest),preserved=preserved,
    delivery_data_unchanged_since_packaging=True,event=state.active_event_id,event_status=state.registry.events[state.active_event_id].status,
    phase_status=[step['status'] for step in steps],report=report.report_id,report_status=report.status,stage_report=report.stage_report,
    pdf_pages=len(pdf),docx_text_length=len(doc_text),pdf_text_length=len(pdf_text),
    files={p.name:sha(p) for p in out.iterdir() if p.is_file()})
(E/'delivery-verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
