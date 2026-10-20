"""Bounded preservation/build/delivery for the approved Qinling simulation."""
from pathlib import Path
from datetime import datetime, timezone
import hashlib, json, shutil, subprocess, sys

W = Path(__file__).resolve().parents[1]
N = W.parent
OLD = N/'release/北山林火连贯演示_20261018_103550'
E = N/'docs/开发/秦岭迁移与离线覆盖证据_20261019'
B = W/'.p00-runs/qinling-20261019'

def sha(p):
    with p.open('rb') as f: return hashlib.file_digest(f, 'sha256').hexdigest()

def inventory(root):
    return {p.relative_to(root).as_posix():sha(p) for p in root.rglob('*')
            if p.is_file() and '__pycache__' not in p.parts and p.suffix!='.pyc'}

def write(p, value):
    p.parent.mkdir(parents=True,exist_ok=True)
    p.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')

def before():
    E.mkdir(); B.mkdir(); (B/'tmp').mkdir()
    value=dict(root=str(OLD),data=inventory(OLD/'data'),source=inventory(OLD/'source'),
               maps=inventory(OLD/'maps'),apk=inventory(OLD/'android'))
    write(E/'before.json',value)
    shutil.copytree(OLD/'data',E/'original-data-before')
    shutil.copytree(OLD/'source',E/'source-before',ignore=shutil.ignore_patterns('settings.yaml','__pycache__','*.pyc'))
    (E/'git-before.diff').write_bytes(subprocess.check_output(['git','diff'],cwd=W))
    (E/'git-before-status.txt').write_bytes(subprocess.check_output(['git','-c','core.quotepath=false','status','--short'],cwd=W))
    for log in (N/'docs/WORKLOG.md',W/'docs/WORKLOG.md'):
        with log.open('a',encoding='utf-8',newline='\n') as f:
            f.write('\n\n### 2026-10-19 秦岭模拟演示迁移与离线高清覆盖盘点（用户同意，实施中）\n'
                    '按已对齐北侧候选33.945638N/108.778381E建立独立秦岭模拟副本，保留原演示数据、源码、地图及审计证据；同步场景几何、三方案、分钟演示、人员物资及复盘。另盘点全包高层级连续片区及坐标；不合并性能修复。\n'
                    '计划code/docs/plans/2026-10-19-qinling-offline-demo.md；证据docs/开发/秦岭迁移与离线覆盖证据_20261019。待处理事项：隔离构建、覆盖核对、当前接口及浏览器验收、独立交付和原数据保全复核。\n')
    print(json.dumps(dict(evidence=str(E),build=str(B)),ensure_ascii=False))

def freeze():
    assert not (B/'process.lock').exists()
    shutil.copytree(W/'impl',B/'source',dirs_exist_ok=True,
        ignore=shutil.ignore_patterns('data','snapshots','__pycache__','*.pyc','settings.yaml'))
    shutil.copy2(OLD/'source/config/settings.yaml',B/'source/config/settings.yaml')
    shutil.copytree(OLD/'guard',B/'guard',dirs_exist_ok=True)
    for name in ('run_fire_demo.py','启动同网手机演示.ps1','启动演示.ps1','停止演示.ps1','安卓App安装与使用.md','手机演示使用说明.md'):
        shutil.copy2(OLD/name,B/name)
    shutil.copytree(OLD/'android',B/'android',dirs_exist_ok=True)
    (B/'maps').mkdir(exist_ok=True)
    shutil.copy2(OLD/'maps/local-satellite.json',B/'maps/local-satellite.json')
    target=B/'maps/google-satellite.sqlite'
    if not target.exists():target.hardlink_to(OLD/'maps/google-satellite.sqlite')
    write(B/'source-manifest.json',inventory(B/'source'))
    runner='''from pathlib import Path
import sys,os,tempfile,json,unittest,time
root=Path(__file__).resolve().parent
sys.dont_write_bytecode=True
for k in list(os.environ):
 if any(v in k.upper() for v in ('KEY','TOKEN','SECRET','PASSWORD','CREDENTIAL')) or k.upper().endswith('_PROXY'):os.environ.pop(k,None)
tempfile.tempdir=str(root/'tmp');os.environ.update(TEMP=str(root/'tmp'),TMP=str(root/'tmp'))
sys.path[:0]=[str(root/'source'),str(root/'guard')]
from p00_baseline import guard
guard(root)
if sys.argv[1]=='build':
 from panda_brain.fire_drill import seed
 state=seed(root/'data',template_id='qinling-north-v1')
 print(json.dumps(dict(event=state.active_event_id,audit=state.workspace.version,valid=not state.workspace.verify_audit_chain()),ensure_ascii=False),flush=True)
elif sys.argv[1]=='test':
 os.environ['PZ_QINLING_ROOT']=str(root)
 suite=unittest.defaultTestLoader.loadTestsFromNames(sys.argv[2:])
 start=time.perf_counter()
 with (root/'tests.log').open('w',encoding='utf-8') as f:r=unittest.TextTestRunner(stream=f,verbosity=2).run(suite)
 value=dict(run=r.testsRun,failures=len(r.failures),errors=len(r.errors),skips=len(r.skipped),seconds=round(time.perf_counter()-start,3),failed=[t.id() for t,_ in r.failures+r.errors])
 (root/'tests.json').write_text(json.dumps(value),encoding='utf-8');print(json.dumps(value),flush=True)
 raise SystemExit(not r.wasSuccessful())
'''
    (B/'verify.py').write_text(runner,encoding='utf-8')
    print(B)

def build():
    assert not (B/'data').exists(),'Never reseed an existing root'
    with (E/'build.log').open('w',encoding='utf-8') as out:
        result=subprocess.run([sys.executable,'-B',str(B/'verify.py'),'build'],stdout=out,stderr=subprocess.STDOUT)
    print(dict(exit=result.returncode,log=str(E/'build.log')))
    raise SystemExit(result.returncode)

def preserve_seed():
    assert json.loads((B/'tests.json').read_text())['failures']==0
    shutil.copytree(B/'data',E/'qinling-seed-data')
    write(E/'qinling-seed-hashes.json',inventory(B/'data'))
    shutil.copy2(B/'tests.log',E/'tests.log');shutil.copy2(B/'tests.json',E/'tests.json')
    print('Independent seed preserved before HTTP/browser QA')

def package():
    import os
    before=json.loads((E/'before.json').read_text(encoding='utf-8'))
    assert inventory(OLD/'data')==before['data'],'Original live data changed; inspect before switching'
    assert inventory(OLD/'source')==before['source']
    assert inventory(OLD/'maps')==before['maps']
    sources=inventory(B/'source')
    assert all(n=='config/settings.yaml' or sha(W/'impl'/n)==h for n,h in sources.items())
    R=N/'release'/('秦岭北侧林火演示_'+datetime.now().strftime('%Y%m%d_%H%M%S'))
    R.mkdir();shutil.copytree(B/'source',R/'source');shutil.copytree(E/'qinling-seed-data',R/'data')
    registry=R/'data/registry.json';value=json.loads(registry.read_text(encoding='utf-8'))
    for event in value['events'].values():
        assert Path(event['snapshot_path']).resolve().is_relative_to((B/'data').resolve())
        event['snapshot_path']=str(R/'data/snapshots'/Path(event['snapshot_path']).name)
    write(registry,value)
    shutil.copytree(B/'guard',R/'guard');shutil.copytree(B/'android',R/'android')
    shutil.copytree(B/'示例表',R/'示例表');(R/'tmp').mkdir();(R/'maps').mkdir()
    shutil.copy2(OLD/'maps/local-satellite.json',R/'maps/local-satellite.json')
    assert shutil.disk_usage(R).free>(OLD/'maps/google-satellite.sqlite').stat().st_size+100*1024*1024
    shutil.copy2(OLD/'maps/google-satellite.sqlite',R/'maps/google-satellite.sqlite')
    assert os.stat(OLD/'maps/google-satellite.sqlite').st_ino!=os.stat(R/'maps/google-satellite.sqlite').st_ino
    for name in ('run_fire_demo.py','启动同网手机演示.ps1','启动演示.ps1','停止演示.ps1','安卓App安装与使用.md','手机演示使用说明.md'):
        shutil.copy2(B/name,R/name)
    write(R/'source-manifest.json',sources)
    instructions='''# 秦岭北侧林火模拟演示

运行本目录“启动同网手机演示.ps1”，电脑沙盘 http://127.0.0.1:18703/sand ，安装页 http://127.0.0.1:18703/wifi 。手机服务地址以启动器显示的当前Wi-Fi地址为准。切换演示后电脑和App均须重新登录，旧会话不沿用。

本演示独立保存。原北山演示及全部原数据/审计/证据保留；两者不能同时占用18703/18705。需要回原演示，先运行本目录“停止演示.ps1”，确认正常退出，再从原目录启动，不能仅删锁或杀PID。

中心33.945638 N、108.778381 E。22项模拟空间资料沿原场景保持近似米制尺度迁移；三方案、分钟重放、两版任务接收执行核验、九步与五步复盘共用新模拟事件。新事件不是原历史迁移，原演示的211条审计仍保存在原目录。

当地卫星影像真实存在，但火场、火线、通道、部署、物资库与外协位置均为虚构演练，未按真实地形、道路或行政归属核验。不得据此认定通道安全、车辆可达或现场发生火灾。

整个约8.3×6.9公里场景外接范围的16级瓦片252/252齐全；更高层级仅局部存在，放大缺图继续明确提示。图像拍摄时间、原始米级分辨率与定位精度未认证。地图在本地，不依赖在线图源；本轮没有合并后台性能修复。

原生安卓0.3.0-demo安装包与原签名不变。模拟副指挥13900000120、值班员13900000119、待到场队长13900000031，演示密码均123456。名单120人、32人已到场、另1队待扫码；饮水120只实收一次，餐食30份、照明4套及下一班看守仍跟进。真实安卓相机、触控及厂商兼容仍待验。
'''
    (R/'开始使用.md').write_text(instructions,encoding='utf-8')
    write(E/'package.json',dict(root=str(R),source_files=len(sources),data=inventory(R/'data'),
        maps=inventory(R/'maps'),apk=inventory(R/'android'),map_independent_from_original=True,
        original_data_preserved=True,original_source_preserved=True))
    (E/'delivery-root.txt').write_text(str(R),encoding='utf-8');print(R)

def stop_verified(root,ports):
    import psutil,time
    root=root.resolve();assert root.is_relative_to(N.resolve())
    runtime=json.loads((root/'runtime.json').read_text(encoding='utf-8'))
    assert Path(runtime['root']).resolve()==root
    proc=psutil.Process(runtime['pid']);command=proc.cmdline()
    assert any(str(root/'run_fire_demo.py').casefold()==arg.casefold() for arg in command)
    assert (root/'process.lock').read_text().strip()==str(proc.pid)
    listening={c.laddr.port for c in proc.net_connections(kind='tcp') if c.status=='LISTEN'}
    assert set(ports)<=listening,(ports,listening)
    started=time.time();(root/'stop').write_text('approved Qinling demonstration switch',encoding='utf-8')
    proc.wait(timeout=30)
    last=root/'last-exit.json';assert last.stat().st_mtime>=started-1
    value=json.loads(last.read_text(encoding='utf-8'))
    assert value['normal'] and value['audit_valid'] and not (root/'process.lock').exists()
    return dict(root=str(root),pid=runtime['pid'],ports=ports,normal_exit=value)

def switch_stop():
    R=Path((E/'delivery-root.txt').read_text(encoding='utf-8'))
    package=json.loads((E/'package.json').read_text(encoding='utf-8'))
    assert inventory(R/'data')==package['data']
    assert inventory(R/'source')==json.loads((R/'source-manifest.json').read_text(encoding='utf-8'))
    stopped=[stop_verified(B,[18704,18706])]
    write(E/'normal-stops.json',stopped)
    before=json.loads((E/'before.json').read_text(encoding='utf-8'))
    assert inventory(OLD/'data')==before['data'] and inventory(OLD/'source')==before['source']
    stopped.append(stop_verified(OLD,[18703,18705]));write(E/'normal-stops.json',stopped)
    assert inventory(OLD/'data')==before['data']
    print(json.dumps(stopped,ensure_ascii=False))

def finish():
    R=Path((E/'delivery-root.txt').read_text(encoding='utf-8'))
    before=json.loads((E/'before.json').read_text(encoding='utf-8'))
    packaged=json.loads((E/'package.json').read_text(encoding='utf-8'))
    actual=inventory(R/'source');expected=json.loads((R/'source-manifest.json').read_text(encoding='utf-8'))
    assert actual==expected
    assert all(n=='config/settings.yaml' or sha(W/'impl'/n)==h for n,h in actual.items())
    assert inventory(OLD/'data')==before['data'] and inventory(OLD/'source')==before['source']
    assert inventory(OLD/'maps')==before['maps'] and inventory(OLD/'android')==before['apk']
    assert inventory(R/'maps')==packaged['maps'] and inventory(R/'android')==packaged['apk']
    data=inventory(R/'data');assert {n:h for n,h in data.items() if n!='registry.json'}=={n:h for n,h in packaged['data'].items() if n!='registry.json'}
    old_registry=json.loads((E/'qinling-seed-data/registry.json').read_text(encoding='utf-8'))
    for event in old_registry['events'].values():event['snapshot_path']=str(R/'data/snapshots'/Path(event['snapshot_path']).name)
    new_registry=json.loads((R/'data/registry.json').read_text(encoding='utf-8'))
    old_auth=old_registry.pop('credentials');new_auth=new_registry.pop('credentials')
    assert old_registry==new_registry and all(new_auth.get(k)==v for k,v in old_auth.items())
    added=set(new_auth)-set(old_auth);assert added<={'13900000031'}
    sys.dont_write_bytecode=True;sys.path.insert(0,str(R/'source'))
    from panda_brain.resource_demo import open_state
    from panda_brain import nine_flow,review_reports,review_exports
    from panda_brain.read_scope import read_scope
    from panda_brain.personnel_groups import group_people
    state=open_state(R/'data')
    with state._event_authority_lock,read_scope():
        assert not state.workspace.verify_audit_chain() and state.workspace.version==210
        steps=nine_flow.summary(state);assert len(steps)==9 and all(not x['stale'] for x in steps)
        report=next(iter(state.review_ledger.reports.values()));assert not review_reports.freshness_error(state,report)
    exports=E/'exports';exports.mkdir(exist_ok=True)
    for fmt in ('md','docx','pdf'):(exports/('秦岭阶段复盘.'+fmt)).write_bytes(review_exports.export_version(report,fmt))
    from docx import Document
    import fitz
    doc=Document(exports/'秦岭阶段复盘.docx')
    doc_text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    pdf=fitz.open(exports/'秦岭阶段复盘.pdf');pdf_text='\n'.join(p.get_text() for p in pdf)
    for needle in ('秦岭','通道','第2版','照明','九步'):
        assert needle in doc_text and needle in pdf_text,needle
    assert '青岚县北山' not in doc_text and '青岚县北山' not in pdf_text
    changed=[n for n,h in actual.items() if before['source'].get(n)!=h]
    result=dict(root=str(R),runtime=json.loads((R/'runtime.json').read_text(encoding='utf-8')),
        source_files=len(actual),source_frozen_W_match=True,changed_source_files=changed,
        original_source_data_maps_apk_preserved=True,original_audit_count=211,
        new_audit_count=state.workspace.version,new_audit_valid=True,new_data_exact_match=data==packaged['data'],
        new_existing_business_data_preserved=True,initialized_demo_credentials=sorted(added),
        scene_objects=22,plans=3,people=len(state.registry.library),
        arrived_squad_members=sum(len(q.arrived_phones) for q in state.ledger.snapshots[state.active_event_id].squads.values()),
        arrived_report_total=sum(p.arrived for p in group_people(state.workspace,state.active_event_id,ledger=state.ledger).values()),
        phase_status=[x['status'] for x in steps],pdf_pages=len(pdf),
        map_independent_from_original=True,apk_unchanged=True,physical_android_tested=False)
    write(E/'verification.json',result)
    (E/'git-after.diff').write_bytes(subprocess.check_output(['git','diff'],cwd=W))
    (E/'git-after-status.txt').write_bytes(subprocess.check_output(['git','-c','core.quotepath=false','status','--short'],cwd=W))
    print(json.dumps(result,ensure_ascii=False))

if __name__=='__main__':
    {'before':before,'freeze':freeze,'build':build,'preserve-seed':preserve_seed,'package':package,'switch-stop':switch_stop,'finish':finish}[sys.argv[1]]()
